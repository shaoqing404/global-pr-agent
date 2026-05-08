from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Iterable, Union

from .legacy.models import PRDraftResult, QualityIssue  # noqa: E402
from .schemas import WorkflowResult


def quality_summary(issues: Iterable[QualityIssue]) -> str:
    issues = list(issues)
    if not issues:
        return "No blocking quality issues found. Human review is still required before external publication."
    lines = []
    for issue in issues:
        lines.append(f"[{issue.severity.upper()}] {issue.dimension}: {issue.message} Suggestion: {issue.suggestion}")
    return "\n".join(lines)


def _markdown_list(items: list[Any]) -> str:
    return "\n".join(f"- {item}" for item in items if item) or "- N/A"


def legacy_result_to_markdown(result: PRDraftResult) -> str:
    issues_md = quality_summary(result.quality_issues)
    titles = "\n".join(f"{idx}. {title}" for idx, title in enumerate(result.title_options, start=1))
    key_messages = "\n".join(f"- {msg}" for msg in result.key_messages)
    channels = "\n".join(f"- {channel}" for channel in result.channel_recommendations)
    return f"""
# 出海品牌 PR 稿件生成结果

## 生成元信息

- Brand: {result.brand.brand_name}
- Product: {result.brand.product_name}
- Target market: {result.brief.target_market}
- Target language: {result.brief.target_language}
- Model used: {result.model_used}
- Fallback mode: {result.fallback_mode}
- Generated at: {result.generated_at}

## 本土化策略

{result.localization_strategy}

## 标题候选

{titles}

## 关键信息

{key_messages}

## PR 正文

{result.press_release}

## 社媒 Caption

{result.social_caption}

## 媒体 Pitch Email

{result.media_pitch_email}

## 翻译与本地化备注

{result.translation_notes}

## 渠道建议

{channels}

## 质量检查

```text
{issues_md}
```
""".strip() + "\n"


def workflow_to_markdown(result: WorkflowResult) -> str:
    plan_lines = []
    for item in result.task_plan.get("task_plan", []):
        plan_lines.append(f"- [{item.get('status', 'pending')}] Step {item.get('step')}: {item.get('name')} -> {item.get('output')}")
    review_lines = []
    for review in result.expert_reviews:
        review_lines.append(
            f"### {review.reviewer} ({review.score}/100, risk: {review.risk_level})\n"
            f"**Role:** {review.role}\n\n"
            f"**Strengths**\n{_markdown_list(review.strengths)}\n\n"
            f"**Issues**\n{_markdown_list(review.issues)}\n\n"
            f"**Suggestions**\n{_markdown_list(review.suggestions)}"
        )
    final_draft = result.revised_draft or result.draft.press_release
    return f"""
# Global PR Agent Output

## 0. Task Metadata

- Brand: {result.metadata.brand}
- Product / Business: {result.metadata.product_or_business}
- Target Market: {result.metadata.target_market}
- Target Language: {result.metadata.target_language}
- PR Objective: {result.metadata.pr_objective}
- Announcement Type: {result.metadata.announcement_type}
- Generated At: {result.generated_at}
- Model Used: {result.draft.model_used}
- Fallback Mode: {result.draft.fallback_mode}

## 1. AI Clarification Questions

{_markdown_list(result.task_plan.get('clarification_questions', []))}

## 2. Agent Task Plan

{chr(10).join(plan_lines)}

## 3. Human Confirmation

- Brand Tone: {result.human_confirmation.brand_tone}
- Product Tone: {result.human_confirmation.product_tone}
- Target Audience: {result.human_confirmation.target_audience}
- Overseas Strategy Baseline: {result.human_confirmation.overseas_strategy_baseline}
- Core PR Angle: {result.human_confirmation.core_pr_angle}
- Localization Strategy: {result.human_confirmation.localization_strategy}

**Risk Constraints**
{_markdown_list(result.human_confirmation.risk_constraints)}

## 4. PR Director Insight

### Brand Insight
{result.insight.brand_insight}

### Market Insight
{result.insight.market_insight}

### PR Narrative
{result.insight.pr_narrative}

### Localization Advice
{result.insight.localization_advice}

### Risk Notes
{_markdown_list(result.insight.risk_notes)}

## 5. Headline Options

{_markdown_list(result.draft.title_options)}

## 6. Press Release Draft

{result.draft.press_release}

## 7. Social Captions

{_markdown_list(result.draft.social_captions)}

## 8. Media Pitch Email

{result.draft.media_pitch_email}

## 9. Chinese Strategy Summary

{result.draft.chinese_strategy_summary}

## 10. Four Specialist Reviews

{chr(10).join(review_lines) if review_lines else 'No expert review has been generated.'}

## 11. Human Feedback and Final Draft

**Human Feedback**

{result.human_feedback or 'No human feedback provided.'}

**Final Draft**

{final_draft}

## 12. Compliance Status

{result.compliance_status}

> Notice: this is a lightweight risk screen for a PR Agent system. It is not legal advice and does not replace real brand, PR, local-market, or legal approval.
""".strip() + "\n"


def result_to_markdown(result: Any) -> str:
    if isinstance(result, WorkflowResult):
        return workflow_to_markdown(result)
    return legacy_result_to_markdown(result)


def save_markdown(result: Any, path: Union[str, Path]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(result_to_markdown(result), encoding="utf-8")
    return path


def save_json(result: Any, path: Union[str, Path]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if hasattr(result, "to_dict"):
        payload = result.to_dict()
    else:
        payload = result
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def save_docx(result: Any, path: Union[str, Path]) -> Path:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("python-docx is required for DOCX export. Run pip install python-docx.") from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    if isinstance(result, WorkflowResult):
        doc.add_heading("Global PR Agent Output", level=1)
        doc.add_paragraph(f"Brand: {result.metadata.brand}")
        doc.add_paragraph(f"Target Market: {result.metadata.target_market}")
        doc.add_paragraph(f"Target Language: {result.metadata.target_language}")
        doc.add_paragraph(f"Generated At: {result.generated_at}")
        doc.add_heading("Task Plan", level=2)
        for item in result.task_plan.get("task_plan", []):
            doc.add_paragraph(f"Step {item.get('step')}: {item.get('name')} - {item.get('status')}", style="List Bullet")
        doc.add_heading("Human Confirmation", level=2)
        doc.add_paragraph(result.human_confirmation.core_pr_angle)
        doc.add_heading("PR Director Insight", level=2)
        doc.add_paragraph(result.insight.pr_narrative)
        doc.add_paragraph(result.insight.market_insight)
        doc.add_heading("Headline Options", level=2)
        for title in result.draft.title_options:
            doc.add_paragraph(title, style="List Number")
        doc.add_heading("Press Release Draft", level=2)
        for paragraph in result.draft.press_release.split("\n\n"):
            text = paragraph.strip()
            if not text:
                continue
            if text.startswith("# "):
                doc.add_heading(text[2:], level=3)
            elif text.startswith("## "):
                doc.add_heading(text[3:], level=3)
            else:
                doc.add_paragraph(text)
        doc.add_heading("Social Media Captions", level=2)
        for caption in result.draft.social_captions:
            doc.add_paragraph(caption, style="List Bullet")
        doc.add_heading("Media Pitch Email", level=2)
        doc.add_paragraph(result.draft.media_pitch_email)
        doc.add_heading("Chinese Strategy Summary", level=2)
        doc.add_paragraph(result.draft.chinese_strategy_summary)
        doc.add_heading("Specialist Reviews", level=2)
        for review in result.expert_reviews:
            doc.add_heading(f"{review.reviewer} - {review.score}/100", level=3)
            doc.add_paragraph(review.role)
            doc.add_paragraph("Suggestions:")
            for item in review.suggestions:
                doc.add_paragraph(item, style="List Bullet")
        doc.add_heading("Compliance Status", level=2)
        doc.add_paragraph(result.compliance_status)
    else:
        # Legacy exporter path.
        doc.add_heading("出海品牌 PR 稿件生成结果", level=1)
        doc.add_paragraph(f"Brand: {result.brand.brand_name}")
        doc.add_paragraph(f"Product: {result.brand.product_name}")
        doc.add_paragraph(f"Target market: {result.brief.target_market}")
        doc.add_heading("PR 正文", level=2)
        doc.add_paragraph(result.press_release)
        doc.add_heading("质量检查", level=2)
        doc.add_paragraph(quality_summary(result.quality_issues))

    doc.save(path)
    return path
